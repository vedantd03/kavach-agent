"""Generate a demo folder that exercises every path the agent has.

    python tools/make_demo_folder.py C:\\kavach\\demo_folder

Every identifier in here is **synthetic**: the Aadhaar numbers carry a valid
Verhoeff check digit and the GSTINs a valid mod-36 one, so the server's
validators actually fire, but the numbers belong to nobody.  Names are made up.

This is the *agent's* corpus: its job is file-type and edge-case coverage
(does chunking, OCR routing and the unscannable path behave?).  It is not the
server's detection eval corpus - that is `kavach-server/corpus/` with its
`labels.jsonl`, owned by the server side (BUILD_PLAN T5).

Folder names are deliberate: Documents, Downloads, Desktop and "OneDrive - Acme"
map to the four interesting `folder_class` values, which drive the exposure
weight in the risk score.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
import zipfile
from pathlib import Path

SEED = 20260926

# --------------------------------------------------------------------------- #
# synthetic identifiers
# --------------------------------------------------------------------------- #

_VERHOEFF_D = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9], [1, 2, 3, 4, 0, 6, 7, 8, 9, 5],
    [2, 3, 4, 0, 1, 7, 8, 9, 5, 6], [3, 4, 0, 1, 2, 8, 9, 5, 6, 7],
    [4, 0, 1, 2, 3, 9, 5, 6, 7, 8], [5, 9, 8, 7, 6, 0, 4, 3, 2, 1],
    [6, 5, 9, 8, 7, 1, 0, 4, 3, 2], [7, 6, 5, 9, 8, 2, 1, 0, 4, 3],
    [8, 7, 6, 5, 9, 3, 2, 1, 0, 4], [9, 8, 7, 6, 5, 4, 3, 2, 1, 0],
]
_VERHOEFF_P = [
    [0, 1, 2, 3, 4, 5, 6, 7, 8, 9], [1, 5, 7, 6, 2, 8, 3, 0, 9, 4],
    [5, 8, 0, 3, 7, 9, 6, 1, 4, 2], [8, 9, 1, 6, 0, 4, 3, 5, 2, 7],
    [9, 4, 5, 3, 1, 2, 6, 8, 7, 0], [4, 2, 8, 6, 5, 7, 3, 9, 0, 1],
    [2, 7, 9, 3, 8, 0, 6, 4, 1, 5], [7, 0, 4, 6, 9, 1, 3, 2, 5, 8],
]
_VERHOEFF_INV = [0, 4, 3, 2, 1, 5, 6, 7, 8, 9]


def verhoeff_digit(number: str) -> str:
    """The check digit that makes ``number + digit`` Verhoeff-valid."""
    check = 0
    for index, char in enumerate(reversed(number)):
        check = _VERHOEFF_D[check][_VERHOEFF_P[(index + 1) % 8][int(char)]]
    return str(_VERHOEFF_INV[check])


def aadhaar(rng: random.Random) -> str:
    """12 digits, first digit 2-9, valid Verhoeff check digit."""
    body = str(rng.randint(2, 9)) + "".join(str(rng.randint(0, 9)) for _ in range(10))
    return body + verhoeff_digit(body)


def spaced(number: str) -> str:
    return f"{number[:4]} {number[4:8]} {number[8:]}"


_PAN_LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def pan(rng: random.Random, holder: str = "P") -> str:
    """AAAPA1234A; the 4th character encodes the holder type (P person, C company)."""
    first = "".join(rng.choice(_PAN_LETTERS) for _ in range(3))
    last_name_initial = rng.choice(_PAN_LETTERS)
    return f"{first}{holder}{last_name_initial}{rng.randint(1000, 9999)}{rng.choice(_PAN_LETTERS)}"


_GSTIN_CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin(rng: random.Random, state: int = 27) -> str:
    """15 characters: state + PAN + entity + Z + mod-36 check character."""
    body = f"{state:02d}{pan(rng, 'C')}{rng.randint(1, 9)}Z"
    total = 0
    for index, char in enumerate(body):
        value = _GSTIN_CHARS.index(char) * (2 if index % 2 else 1)
        total += value // 36 + value % 36
    return body + _GSTIN_CHARS[(36 - total % 36) % 36]


def mobile(rng: random.Random) -> str:
    return f"{rng.choice('6789')}{''.join(str(rng.randint(0, 9)) for _ in range(9))}"


def ifsc(rng: random.Random) -> str:
    return f"{''.join(rng.choice(_PAN_LETTERS) for _ in range(4))}0{rng.randint(100000, 999999)}"


def account(rng: random.Random) -> str:
    return "".join(str(rng.randint(0, 9)) for _ in range(rng.choice([11, 12, 14])))


FIRST = ["Ravi", "Priya", "Arjun", "Meera", "Sanjay", "Anita", "Vikram", "Deepa", "Rahul", "Kavya"]
LAST = ["Kumar", "Sharma", "Iyer", "Patel", "Nair", "Reddy", "Gupta", "Bose", "Menon", "Joshi"]


def person(rng: random.Random) -> tuple[str, str]:
    name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
    email = name.lower().replace(" ", ".") + f"{rng.randint(1, 99)}@example.in"
    return name, email


# --------------------------------------------------------------------------- #
# builders
# --------------------------------------------------------------------------- #

def write(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def build_text_files(root: Path, rng: random.Random) -> None:
    docs, downloads, desktop = root / "Documents", root / "Downloads", root / "Desktop"

    name, email = person(rng)
    write(docs / "whatsapp_chat_support.txt",
          f"[12/09/26, 10:14] {name.split()[0]}: bhai mera aadhar no hai {spaced(aadhaar(rng))}\n"
          f"[12/09/26, 10:15] Support: PAN bhi bhej dijiye\n"
          f"[12/09/26, 10:16] {name.split()[0]}: {pan(rng)}, mobile {mobile(rng)}\n"
          f"[12/09/26, 10:18] Support: mail id {email} pe bhej diya hai\n")

    # Devanagari: the lexicon carries Hindi cues, so the demo should too.
    write(docs / "kyc_note_hindi.txt",
          f"ग्राहक का नाम: {person(rng)[0]}\n"
          f"आधार संख्या: {spaced(aadhaar(rng))}\n"
          f"पैन कार्ड: {pan(rng)}\n"
          f"मोबाइल: {mobile(rng)}\n")

    write(docs / "notes.md",
          "# Sprint notes\n\n- ship the detector\n- no personal data in this file\n"
          "- follow up with the vendor about the renewal\n")

    write(desktop / "salary_slip.txt",
          f"Employee: {person(rng)[0]}\nPAN: {pan(rng)}\n"
          f"Bank A/C: {account(rng)}\nIFSC: {ifsc(rng)}\nNet pay: 84,250\n")

    write(downloads / "app_config.json", json.dumps({
        "service": "billing", "region": "ap-south-1",
        "support_email": person(rng)[1], "contact": mobile(rng),
        "gstin": gstin(rng), "retries": 3,
    }, indent=2))

    write(downloads / "server.log", "".join(
        f"2026-09-2{i % 9} 11:0{i % 6}:00 INFO  request user={person(rng)[1]} "
        f"phone={mobile(rng)} status=200\n" for i in range(40)))

    write(root / ".env",
          "DB_HOST=db.internal\nDB_USER=billing\n"
          "AWS_SECRET_ACCESS_KEY=wJalrXUtnFEMIKEXAMPLEKEYnotreal1234567890\n"
          "STRIPE_KEY=sk_live_notarealkeyjustfortesting0000\n")

    write(downloads / "settings.ini",
          "[smtp]\nhost = mail.example.in\nuser = billing@example.in\n"
          f"[oncall]\nphone = {mobile(rng)}\n")

    write(downloads / "vendor.yaml",
          f"vendor:\n  name: Acme Supplies\n  gstin: {gstin(rng, 29)}\n"
          f"  pan: {pan(rng, 'C')}\n  contact: {mobile(rng)}\n")
    write(downloads / "pipeline.yml", "stages:\n  - build\n  - test\n  - deploy\n")

    write(docs / "invoice_meta.xml",
          "<invoice>\n  <number>INV-2026-0912</number>\n"
          f"  <buyer_gstin>{gstin(rng)}</buyer_gstin>\n"
          f"  <contact>{mobile(rng)}</contact>\n</invoice>\n")

    buyer, buyer_email = person(rng)
    write(docs / "order_confirmation.eml",
          f"From: orders@example.in\nTo: {buyer_email}\nSubject: Order confirmed\n\n"
          f"Hi {buyer.split()[0]},\nYour order ships tomorrow. "
          f"We have your number as {mobile(rng)}.\n")

    # Hard negatives: 12-digit runs that must NOT be read as Aadhaar.
    write(downloads / "logistics_tracking.txt",
          "Order ID 481920571034 dispatched\nAWB 998877665544 in transit\n"
          "UTR 100234998877 settled\nInvoice no 2026091200412\n"
          "Reference 774411223399 closed\n")

    # Already masked in the source: should be flagged masked_in_source, low risk.
    write(docs / "masked_extract.txt",
          "Verified customers (masked export)\n"
          "Aadhaar XXXXXXXX4821\nPAN XXXXX1234F\nMobile XXXXXX3210\n")

    write(docs / "empty.txt", "")


def build_sheets(root: Path, rng: random.Random) -> None:
    downloads = root / "Downloads"
    downloads.mkdir(parents=True, exist_ok=True)

    rows = []
    for _ in range(120):                       # 3 blocks of 50 -> multi-block chunking
        name, email = person(rng)
        rows.append((name, email, mobile(rng), aadhaar(rng)))
    write(downloads / "customer_export_aug.csv",
          "name,email,mobile,aadhaar\n" + "".join(",".join(r) + "\n" for r in rows))

    write(downloads / "partners.tsv",
          "partner\tgstin\tcontact\n" + "".join(
              f"{person(rng)[0]}\t{gstin(rng)}\t{mobile(rng)}\n" for _ in range(8)))

    try:
        import openpyxl
    except ImportError:
        print("  ! openpyxl missing, skipping payroll.xlsx")
        return
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "payroll"
    sheet.append(["employee", "pan", "account", "ifsc", "net_pay"])
    for _ in range(60):
        sheet.append([person(rng)[0], pan(rng), account(rng), ifsc(rng), rng.randint(40000, 150000)])
    summary = workbook.create_sheet("summary")      # second sheet -> page=2
    summary.append(["metric", "value"])
    summary.append(["headcount", 60])
    workbook.create_sheet("blank")                  # empty sheet -> no chunks
    workbook.save(downloads / "payroll.xlsx")


def build_docx(root: Path, rng: random.Random) -> None:
    try:
        import docx
    except ImportError:
        print("  ! python-docx missing, skipping kyc_record.docx")
        return
    document = docx.Document()
    document.add_paragraph("Employee KYC record")
    document.add_paragraph("STRICTLY CONFIDENTIAL")       # marking -> tier should rise
    name, email = person(rng)
    document.add_paragraph(f"Name: {name}")
    document.add_paragraph(f"Aadhaar: {spaced(aadhaar(rng))}")
    table = document.add_table(rows=3, cols=2)
    for row, (key, value) in enumerate(
        [("PAN", pan(rng)), ("Mobile", mobile(rng)), ("Email", email)]
    ):
        table.cell(row, 0).text = key
        table.cell(row, 1).text = value
    document.save(root / "Documents" / "kyc_record.docx")


def build_pdfs(root: Path, rng: random.Random) -> None:
    docs = root / "Documents"
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        print("  ! reportlab missing, skipping PDFs  (pip install reportlab)")
        return

    pdf = canvas.Canvas(str(docs / "invoice.pdf"), pagesize=A4)
    pdf.drawString(72, 780, "TAX INVOICE  -  Acme Supplies Pvt Ltd")
    pdf.drawString(72, 760, f"GSTIN {gstin(rng)}   PAN {pan(rng, 'C')}")
    pdf.drawString(72, 740, "Invoice no INV-2026-0912   Date 12 Sep 2026")
    pdf.drawString(72, 720, f"Buyer contact {mobile(rng)}")
    pdf.showPage()
    pdf.drawString(72, 780, "Page two carries plenty of extractable text as well.")
    pdf.showPage()
    pdf.drawString(72, 780, ".")            # under 20 chars -> treated as scanned -> /ocr
    pdf.showPage()
    pdf.save()

    broken = docs / "corrupt.pdf"
    broken.write_bytes(b"%PDF-1.4\nthis file is truncated and cannot be parsed\n")


def build_images(root: Path, rng: random.Random) -> None:
    docs = root / "Documents"
    try:
        from PIL import Image, ImageDraw
    except ImportError:
        print("  ! Pillow missing, skipping images")
        return

    def card(path: Path, lines: list[str], fmt: str | None = None) -> None:
        image = Image.new("RGB", (620, 200), "white")
        draw = ImageDraw.Draw(image)
        for index, line in enumerate(lines):
            draw.text((20, 30 + index * 30), line, fill="black")
        image.save(path, format=fmt)

    name = person(rng)[0]
    card(docs / "kyc_scan.png", ["GOVERNMENT OF INDIA", name, f"Aadhaar {spaced(aadhaar(rng))}"])
    card(docs / "pan_card.jpg", ["INCOME TAX DEPARTMENT", name, f"PAN {pan(rng)}"])
    card(docs / "cheque.tiff", ["STATE BANK", f"A/C {account(rng)}", f"IFSC {ifsc(rng)}"], "TIFF")


def build_edge_cases(root: Path, rng: random.Random) -> None:
    """Files that must come back unscannable, skipped or ignored."""
    downloads, synced = root / "Downloads", root / "OneDrive - Acme"
    shared = root / "Shared"

    # OLE magic = a password-protected Office file -> status_reason "encrypted"
    (downloads / "locked.docx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 1024)

    # A zip that claims to be a workbook -> "corrupt"
    broken = downloads / "broken.xlsx"
    with zipfile.ZipFile(broken, "w") as archive:
        archive.writestr("not-a-workbook.txt", "nope")

    # folder_class coverage: synced and shared weigh highest in the risk score
    name, email = person(rng)
    write(synced / "client_list.csv",
          "client,email,pan\n" + "".join(
              f"{person(rng)[0]},{person(rng)[1]},{pan(rng, 'C')}\n" for _ in range(12)))
    write(shared / "team_contacts.txt",
          "".join(f"{person(rng)[0]} {mobile(rng)} {person(rng)[1]}\n" for _ in range(10)))

    # Excluded directories and non-indexed extensions: must never be opened.
    write(root / "node_modules" / "pkg" / "index.txt", f"aadhaar {spaced(aadhaar(rng))}")
    write(root / ".git" / "COMMIT_EDITMSG", f"contact {mobile(rng)}")
    write(root / ".venv" / "lib" / "site.txt", "should be excluded")
    (root / "archive.zip").write_bytes(b"PK\x03\x04 not scanned, extension not included")
    (root / "photo.heic").write_bytes(b"\x00\x00\x00 not an indexed type")


MANIFEST = """# Demo folder - what the agent should do with each file

Generated by `tools/make_demo_folder.py`. Every identifier is synthetic; the
Aadhaar and GSTIN check digits are valid so the server's validators fire.

## Expected agent behaviour

| Path | Type | Expected |
|---|---|---|
| Documents/whatsapp_chat_support.txt | txt | Hinglish cues, Aadhaar + PAN + mobile + email |
| Documents/kyc_note_hindi.txt | txt | Devanagari cues and digits |
| Documents/notes.md | md | no findings |
| Documents/masked_extract.txt | txt | `masked_in_source`, low risk |
| Documents/empty.txt | txt | status ok, **zero chunks** |
| Documents/invoice_meta.xml | xml | GSTIN, mobile |
| Documents/order_confirmation.eml | eml | email + mobile |
| Documents/kyc_record.docx | docx | paragraphs + table rows; "STRICTLY CONFIDENTIAL" marking |
| Documents/invoice.pdf | pdf | pages 1-2 text, page 3 rendered to `/ocr` |
| Documents/corrupt.pdf | pdf | **unscannable**, `corrupt` |
| Documents/kyc_scan.png | png | `/ocr`, page=1, `ocr_confidence` set |
| Documents/pan_card.jpg | jpg | `/ocr` as image/jpeg |
| Documents/cheque.tiff | tiff | first frame converted to PNG, then `/ocr` |
| Desktop/salary_slip.txt | txt | PAN + account + IFSC, `folder_class=desktop` |
| Downloads/customer_export_aug.csv | csv | **column chunks**, 3 blocks (rows 2, 52, 102); bulk -> restricted |
| Downloads/partners.tsv | tsv | tab-delimited column chunks |
| Downloads/payroll.xlsx | xlsx | sheet 1 + sheet 2 (`page`), blank sheet yields nothing |
| Downloads/app_config.json | json | email, mobile, GSTIN |
| Downloads/server.log | log | repeated emails and mobiles across chunks |
| Downloads/settings.ini | ini | contact details |
| Downloads/vendor.yaml | yaml | business PII |
| Downloads/pipeline.yml | yml | no findings |
| Downloads/logistics_tracking.txt | txt | **hard negatives** - 12-digit ids that are not Aadhaar |
| Downloads/locked.docx | docx | **unscannable**, `encrypted` |
| Downloads/broken.xlsx | xlsx | **unscannable**, `corrupt` |
| .env | env | secrets (AWS/Stripe shaped, not real) |
| OneDrive - Acme/client_list.csv | csv | `folder_class=synced`, highest exposure weight |
| Shared/team_contacts.txt | txt | `folder_class=shared` |
| node_modules/, .git/, .venv/ | - | **never opened** (excluded dirs) |
| archive.zip, photo.heic | - | **ignored**, extension not in include_types |

## Not included by default

An oversize file, because it would need to exceed `max_file_mb` (50 MB).
Test that path instead by lowering the limit for one run:

    MAX_FILE_MB=1 kavach-agent.exe scan <this folder>

Any file over 1 MB then comes back `unscannable` / `oversize` without being read.
"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("target", nargs="?", default="demo_folder", help="folder to create")
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args()

    root = Path(args.target).expanduser().resolve()
    if root.exists() and any(root.iterdir()):
        print(f"refusing to write into a non-empty folder: {root}", file=sys.stderr)
        print("delete it first, or pass a different path", file=sys.stderr)
        return 2

    rng = random.Random(args.seed)
    root.mkdir(parents=True, exist_ok=True)
    print(f"generating demo corpus in {root}")
    build_text_files(root, rng)
    build_sheets(root, rng)
    build_docx(root, rng)
    build_pdfs(root, rng)
    build_images(root, rng)
    build_edge_cases(root, rng)
    (root / "MANIFEST.md").write_text(MANIFEST, encoding="utf-8")

    files = sorted(p for p in root.rglob("*") if p.is_file())
    total = sum(p.stat().st_size for p in files)
    print(f"  {len(files)} files, {total / 1024:.0f} KB")
    for path in files:
        print(f"    {path.relative_to(root).as_posix():<44} {path.stat().st_size:>8} B")
    print("\nscan it with:")
    print(f"    python -m agent.main scan \"{root}\"")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
